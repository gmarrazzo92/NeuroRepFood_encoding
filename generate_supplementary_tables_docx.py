#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
Generate a Word document containing the canonical supplementary tables.

The script reads the manuscript-facing CSVs produced by the analysis pipeline
and writes them, in canonical S-table order, to a single .docx file.

It does NOT add scientific captions, figure panels, or explanatory prose.
Only a simple table identifier (e.g. "Supplementary Table S5b") is inserted
before each table so that the tables are easy to locate while assembling the
final Supplementary Material.

Default repository layout
-------------------------
<repository>/
    reproduced_outputs/
        manuscript/
            tables/
                Supplementary_Table_S1.csv
                ...
    generate_supplementary_tables_docx.py

Usage
-----
From the repository root:

    python generate_supplementary_tables_docx.py

Explicit paths:

    python generate_supplementary_tables_docx.py ^
        --tables-dir reproduced_outputs\\manuscript\\tables ^
        --output Supplementary_Tables.docx

The CSV cell contents are transferred as strings without numerical
recalculation or re-rounding.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


# Canonical manuscript-table order.
TABLE_ORDER = [
    "S1", "S2", "S3", "S4", "S4b",
    "S5", "S5b", "S5c", "S5d",
    "S6", "S7", "S7b",
    "S8", "S8b", "S9", "S10",
    "S11", "S12", "S13", "S14",
]

A4_W = Mm(210)
A4_H = Mm(297)

# Tables with 7+ columns are placed in landscape by default.
LANDSCAPE_COLUMN_THRESHOLD = 7

FONT_NAME = "Arial"
BODY_FONT_SIZE = Pt(8)
WIDE_TABLE_FONT_SIZE = Pt(7.5)
HEADER_FONT_SIZE = Pt(8)
WIDE_HEADER_FONT_SIZE = Pt(7.5)
LABEL_FONT_SIZE = Pt(11)

MARGIN_PORTRAIT = Mm(15)
MARGIN_LANDSCAPE = Mm(12)


def set_run_font(run, size: Pt, bold: bool | None = None):
    """Set font robustly for Word/LibreOffice."""
    run.font.name = FONT_NAME
    run.font.size = size
    if bold is not None:
        run.bold = bold

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), FONT_NAME)
    rfonts.set(qn("w:hAnsi"), FONT_NAME)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)


def set_section_layout(section, landscape: bool):
    """Set A4 orientation and compact margins."""
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = A4_H
        section.page_height = A4_W
        margin = MARGIN_LANDSCAPE
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = A4_W
        section.page_height = A4_H
        margin = MARGIN_PORTRAIT

    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin


def set_repeat_table_header(row):
    """Mark the first table row to repeat on subsequent pages."""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    """Prevent a table row from splitting across pages where possible."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_margins(cell, top=55, start=70, bottom=55, end=70):
    """
    Set compact cell margins in twips.
    72 twips ~= 0.05 inches.
    """
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for tag, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill="E7E6E6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, size="4", color="B7B7B7"):
    """Apply thin neutral borders to all table edges/internal rules."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def visible_length(value: str) -> int:
    """Approximate display length for proportional column allocation."""
    value = str(value)
    return max((len(line) for line in value.splitlines()), default=1)


def column_widths(df: pd.DataFrame, usable_width_cm: float) -> list[float]:
    """
    Allocate widths using observed text length, with conservative caps.
    This is a layout heuristic only; cell content itself is untouched.
    """
    raw = []
    for col in df.columns:
        vals = [str(col)] + df[col].astype(str).tolist()
        max_len = max(visible_length(v) for v in vals)

        # Compress influence of very long labels so one text column does not
        # consume the entire page.
        weight = max(4.0, min(max_len, 34) ** 0.65)
        raw.append(weight)

    total = sum(raw)
    widths = [usable_width_cm * w / total for w in raw]

    # Guarantee a practical minimum; re-normalize if needed.
    min_w = 1.15 if len(widths) >= 8 else 1.35
    widths = [max(min_w, w) for w in widths]
    scale = usable_width_cm / sum(widths)
    return [w * scale for w in widths]


def write_cell_text(cell, value: str, *, header: bool, wide: bool, first_col: bool):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT if first_col else WD_ALIGN_PARAGRAPH.CENTER
    )

    run = p.add_run(str(value))
    size = (
        WIDE_HEADER_FONT_SIZE if (wide and header)
        else HEADER_FONT_SIZE if header
        else WIDE_TABLE_FONT_SIZE if wide
        else BODY_FONT_SIZE
    )
    set_run_font(run, size=size, bold=header)

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)

    if header:
        shade_cell(cell)


def add_dataframe_table(doc: Document, df: pd.DataFrame, section, wide: bool):
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    set_table_borders(table)

    # Header.
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for j, col in enumerate(df.columns):
        write_cell_text(
            hdr.cells[j], col, header=True, wide=wide, first_col=(j == 0)
        )

    # Body.
    for _, row in df.iterrows():
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for j, value in enumerate(row.tolist()):
            write_cell_text(
                cells[j],
                value,
                header=False,
                wide=wide,
                first_col=(j == 0),
            )

    usable_cm = (
        (section.page_width - section.left_margin - section.right_margin) / Cm(1)
    )
    widths = column_widths(df, float(usable_cm))

    for j, width_cm in enumerate(widths):
        width = Cm(width_cm)
        for row in table.rows:
            row.cells[j].width = width

    return table


def table_id_from_filename(path: Path) -> str:
    m = re.fullmatch(r"Supplementary_Table_(S\d+[a-z]?)\.csv", path.name)
    if not m:
        raise ValueError(f"Unexpected table filename: {path.name}")
    return m.group(1)


def load_tables(tables_dir: Path, allow_missing: bool):
    expected = {
        tid: tables_dir / f"Supplementary_Table_{tid}.csv"
        for tid in TABLE_ORDER
    }

    missing = [str(path) for path in expected.values() if not path.is_file()]
    if missing and not allow_missing:
        raise FileNotFoundError(
            "Missing canonical manuscript table(s):\n  " + "\n  ".join(missing)
        )

    found_extra = []
    for path in tables_dir.glob("Supplementary_Table_S*.csv"):
        try:
            tid = table_id_from_filename(path)
        except ValueError:
            continue
        if tid not in TABLE_ORDER:
            found_extra.append(path.name)

    if found_extra:
        print(
            "NOTE: additional Supplementary_Table_*.csv files were found but are "
            "not part of the canonical manuscript table order and will not be "
            f"included: {', '.join(sorted(found_extra))}"
        )

    tables = []
    for tid in TABLE_ORDER:
        path = expected[tid]
        if not path.is_file():
            continue

        # Preserve manuscript-facing strings exactly as written by the
        # reproduction pipeline. No numerical parsing/reformatting.
        df = pd.read_csv(path, dtype=str, keep_default_na=False)
        tables.append((tid, path, df))

    if not tables:
        raise FileNotFoundError(f"No canonical supplementary-table CSVs in {tables_dir}")

    return tables


def build_document(tables, output: Path, orientation_mode: str):
    doc = Document()

    # Remove the default empty paragraph only if it remains empty.
    if len(doc.paragraphs) == 1 and not doc.paragraphs[0].text:
        p = doc.paragraphs[0]._element
        p.getparent().remove(p)

    for i, (tid, source, df) in enumerate(tables):
        if orientation_mode == "landscape":
            wide = True
        elif orientation_mode == "portrait":
            wide = False
        else:
            wide = len(df.columns) >= LANDSCAPE_COLUMN_THRESHOLD

        if i == 0:
            section = doc.sections[0]
        else:
            section = doc.add_section(WD_SECTION.NEW_PAGE)

        set_section_layout(section, landscape=wide)

        # Identifier only; the user can replace/extend this with the final
        # scientific caption while assembling the Supplementary Material.
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(f"Supplementary Table {tid}")
        set_run_font(run, LABEL_FONT_SIZE, bold=True)

        add_dataframe_table(doc, df, section, wide=wide)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def parse_args():
    repo_root = Path(__file__).resolve().parent
    default_tables = repo_root / "reproduced_outputs" / "manuscript" / "tables"
    default_output = (
        repo_root
        / "reproduced_outputs"
        / "manuscript"
        / "Supplementary_Tables.docx"
    )

    p = argparse.ArgumentParser(
        description=(
            "Generate a Word document containing all canonical supplementary "
            "tables from the manuscript-facing CSV outputs."
        )
    )
    p.add_argument(
        "--tables-dir",
        type=Path,
        default=default_tables,
        help=f"Directory containing Supplementary_Table_*.csv (default: {default_tables})",
    )
    p.add_argument(
        "--output",
        type=Path,
        default=default_output,
        help=f"Output .docx path (default: {default_output})",
    )
    p.add_argument(
        "--orientation",
        choices=["auto", "portrait", "landscape"],
        default="auto",
        help=(
            "Page orientation. 'auto' uses landscape for tables with 7+ columns "
            "and portrait otherwise."
        ),
    )
    p.add_argument(
        "--allow-missing",
        action="store_true",
        help="Generate from the available canonical tables instead of failing if one is missing.",
    )
    return p.parse_args()


def main():
    args = parse_args()

    tables_dir = args.tables_dir.resolve()
    output = args.output.resolve()

    print("=" * 88)
    print("GENERATE SUPPLEMENTARY TABLES DOCX")
    print("=" * 88)
    print(f"Tables : {tables_dir}")
    print(f"Output : {output}")

    tables = load_tables(tables_dir, allow_missing=args.allow_missing)

    print(f"\nFound {len(tables)} canonical tables:")
    for tid, path, df in tables:
        print(f"  {tid:4s}  {df.shape[0]:2d} rows x {df.shape[1]:2d} cols  {path.name}")

    build_document(
        tables=tables,
        output=output,
        orientation_mode=args.orientation,
    )

    print(f"\nSaved: {output}")


if __name__ == "__main__":
    main()
